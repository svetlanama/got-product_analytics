# -*- coding: utf-8 -*-
"""Build the document deliverables of the final project:
  Task 1 — A/B test plan                      (…_Final Project-1.docx)
  Task 2 — tracking plan                      (…_Final Project-2.pdf, via HTML + headless Chrome)
  Task 3 — statistical results & interpretation (…_Final Project-3.docx)
Needs python-docx, pandas, scipy, matplotlib; Google Chrome for the PDF."""
import subprocess

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from fp_common import (C_CONTROL, C_GRID, C_INK, C_MUTED, C_TEST, FIG, HERE, METRIC_UA, STEPS, STEP_UA, SURNAME,
                       add_daily_metrics, daily_tests, load_all, ratio_table, sample_size_two_prop, totals, two_prop)

FIG.mkdir(exist_ok=True)
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
DATE = "17 жовтня 2026 р."
AUTHOR = "Мойсеєнко Світлана"

# ----------------------------------------------------------------------------- analysis
df = load_all()
d = add_daily_metrics(df[df.complete])
tot = totals(df[df.complete])
ratios = ratio_table(tot)
tests = daily_tests(d, ["spend", "impressions", "reach", "clicks", "searches", "view_content", "add_to_cart",
                        "purchase", "ctr", "cpc", "cart_rate", "cr_cart", "cr_click", "cpa"])
T, C = d[d.group == "Test"], d[d.group == "Control"]
prop_cart = two_prop(T.purchase.sum(), T.add_to_cart.sum(), C.purchase.sum(), C.add_to_cart.sum())
prop_click = two_prop(T.purchase.sum(), T.clicks.sum(), C.purchase.sum(), C.clicks.sum())
prop_ctr = two_prop(T.clicks.sum(), T.impressions.sum(), C.clicks.sum(), C.impressions.sum())
prop_cartrate = two_prop(T.add_to_cart.sum(), T.clicks.sum(), C.add_to_cart.sum(), C.clicks.sum())
anom = d[d.purchase > d.add_to_cart]

# sample size for the test plan (baseline click→purchase ≈ 10 %, relative MDE 10 %)
BASE_CR, MDE_REL = 0.10, 0.10
n_clicks_needed = sample_size_two_prop(BASE_CR, BASE_CR * (1 + MDE_REL))
clicks_per_day = int(C.clicks.mean())
days_needed = int(np.ceil(n_clicks_needed / clicks_per_day))
# MDE for day-level CPA comparison with 30 days per arm (Welch, power 0.8)
from scipy import stats as _st
sd_cpa = C.cpa.std(ddof=1)
mde_cpa = (_st.norm.ppf(0.975) + _st.norm.ppf(0.8)) * sd_cpa * np.sqrt(2 / 30)

# ----------------------------------------------------------------------------- figures
plt.rcParams.update({
    "figure.dpi": 150, "axes.spines.top": False, "axes.spines.right": False, "axes.edgecolor": "#c3c2b7",
    "axes.labelcolor": C_INK, "xtick.color": C_MUTED, "ytick.color": C_MUTED, "axes.grid": True,
    "grid.color": C_GRID, "grid.linewidth": 0.8, "axes.axisbelow": True, "font.size": 9,
})


def fig_daily():
    fig, axes = plt.subplots(1, 3, figsize=(10.5, 3.1))
    panels = [("purchase", "Покупки / день", "{:.0f}"), ("cpa", "CPA, $ за покупку", "${:.1f}"),
              ("cr_cart", "Конверсія кошик → покупка", "{:.0%}")]
    for ax, (m, title, fmt) in zip(axes, panels):
        for grp, color in [("Control", C_CONTROL), ("Test", C_TEST)]:
            s = d[d.group == grp].set_index("date")[m]
            ax.plot(s.index, s.values, color=color, lw=1.8, label=grp)
        ax.set_title(title, loc="left", color=C_INK, fontsize=9.5)
        ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _, f=fmt: f.format(v)))
        ax.grid(axis="x", visible=False); ax.tick_params(axis="x", rotation=45, labelsize=7.5)
    axes[0].legend(frameon=False, loc="lower left", ncol=2, fontsize=8)
    plt.tight_layout(); p = FIG / "daily_metrics.png"; fig.savefig(p); plt.close(fig); return p


def fig_effects():
    rows = [("purchase", "Покупки / день"), ("cpa", "CPA"), ("cr_click", "Клік → покупка"),
            ("cr_cart", "Кошик → покупка"), ("cart_rate", "Клік → кошик"), ("ctr", "CTR"), ("spend", "Витрати / день")]
    fig, ax = plt.subplots(figsize=(7.2, 3.3))
    for i, (m, label) in enumerate(rows[::-1]):
        r = tests.loc[m]
        lo, hi = r.ci_lo / r.Control, r.ci_hi / r.Control
        color = C_TEST if r.p_welch < 0.05 else C_MUTED
        ax.plot([lo, hi], [i, i], color=color, lw=2.2); ax.plot(r.rel, i, "o", color=color, ms=6)
        p_txt = "p < 0.001" if r.p_welch < 0.001 else f"p = {r.p_welch:.3f}"
        ax.text(max(hi, 0) + 0.04, i, f"{r.rel:+.0%}  ({p_txt})", va="center", fontsize=8, color=C_INK)
    ax.axvline(0, color=C_INK, lw=1, ls="--")
    ax.set_yticks(range(len(rows))); ax.set_yticklabels([l for _, l in rows[::-1]], color=C_INK)
    ax.xaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:+.0%}"))
    ax.set_xlim(-0.75, 1.9); ax.grid(axis="y", visible=False)
    ax.set_xlabel("Відносна різниця Test vs Control (середні по днях), 95% ДІ Уелча")
    ax.set_title("Ефект нового підходу за ключовими метриками (синім — значущі при α = 0.05)", loc="left",
                 color=C_INK, fontsize=9.5)
    plt.tight_layout(); p = FIG / "effects.png"; fig.savefig(p); plt.close(fig); return p


P_DAILY, P_EFFECTS = fig_daily(), fig_effects()

# ----------------------------------------------------------------------------- docx helpers
FONT = "Arial"


def new_doc(title, subtitle):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2.2)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
    for name, size, color in [("Heading 1", 16, "1C5CAB"), ("Heading 2", 13, "1C5CAB"), ("Heading 3", 11, "0B0B0B")]:
        h = doc.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = RGBColor.from_string(color); h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        h.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 0); h.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph(); r = p.add_run(title); r.bold = True; r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string("0B0B0B")
    p = doc.add_paragraph(); r = p.add_run(subtitle); r.font.size = Pt(10.5); r.font.color.rgb = RGBColor.from_string("52514E")
    p = doc.add_paragraph(); r = p.add_run(f"Автор: {AUTHOR}  ·  Дисципліна: Продуктова аналітика та прикладна статистика "
                                           f"(GoIT Neoversity)  ·  {DATE}")
    r.font.size = Pt(9.5); r.font.color.rgb = RGBColor.from_string("52514E")
    _bottom_border(p)
    return doc


def _bottom_border(p):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in [("w:val", "single"), ("w:sz", "6"), ("w:space", "4"), ("w:color", "C3C2B7")]:
        bt.set(qn(k), v)
    b.append(bt); pPr.append(b)


def h(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text, bold_lead=None, italic=False, size=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True
    r = p.add_run(text); r.italic = italic
    if size: r.font.size = Pt(size)
    return p


def bullets(doc, items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            r = p.add_run(it[0]); r.bold = True; p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(3)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color); tcPr.append(shd)


def table(doc, header, rows, widths=None, font_size=9, align_right_from=1, zebra=True):
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(text); r.bold = True; r.font.size = Pt(font_size)
        shade(c, "E8EEF7")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""; p = cells[i].paragraphs[0]; r = p.add_run(str(val)); r.font.size = Pt(font_size)
            if align_right_from is not None and i >= align_right_from: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if zebra and ri % 2 == 1: shade(cells[i], "F7F7F4")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def kv_table(doc, pairs, widths=(5.0, 11.6), font_size=9.5):
    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, v in pairs:
        cells = t.add_row().cells
        cells[0].text = ""; r = cells[0].paragraphs[0].add_run(k); r.bold = True; r.font.size = Pt(font_size); shade(cells[0], "F1F0EA")
        cells[1].text = ""; r = cells[1].paragraphs[0].add_run(v); r.font.size = Pt(font_size)
    for row in t.rows:
        for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def fmt_money(v): return f"${v:,.0f}".replace(",", " ")
def fmt_int(v): return f"{v:,.0f}".replace(",", " ")
def fmt_pct(v, nd=1): return f"{v*100:.{nd}f} %"
def fmt_pp(v, nd=1): return f"{v*100:+.{nd}f} п.п."
def fmt_p(p): return "< 0.001" if p < 0.001 else f"{p:.3f}"
def ptxt(p): return "p < 0.001" if p < 0.001 else f"p = {p:.3f}"


# ============================================================================= TASK 1 — TEST PLAN
def build_task1():
    doc = new_doc("Тест-план A/B-експерименту",
                  "Новий підхід до розміщення онлайн-реклами продукту · Фінальний проєкт, Завдання №1")

    h(doc, "1. Контекст і проблема")
    para(doc, "Команда маркетингу пропонує новий підхід до розміщення онлайн-реклами продукту (інші майданчики, "
              "таргетинг та формати оголошень). Поточна кампанія стабільно приводить покупців, але її ефективність "
              "(вартість залучення покупця) зростає. Перш ніж перерозподіляти бюджет, потрібно переконатися, що новий "
              "підхід дає не гірший, а бажано кращий результат за ті самі гроші. Для цього проводимо контрольований "
              "експеримент: контрольна група бачить рекламу як раніше, тестова — за новим підходом.")

    h(doc, "2. Гіпотези")
    kv_table(doc, [
        ("Бізнес-гіпотеза", "Новий підхід до розміщення реклами приводить на сайт більш релевантну аудиторію, "
                            "тому за той самий бюджет ми отримуємо більше покупок, а вартість однієї покупки (CPA) знижується."),
        ("H0 (нульова)", "Середня вартість покупки (CPA) та кількість покупок на день у тестовій групі не відрізняються від контрольної."),
        ("H1 (альтернативна)", "CPA у тестовій групі нижчий, ніж у контрольній (або кількість покупок на день вища) за статистично "
                               "значущої різниці на рівні α = 0.05."),
        ("Напрям тесту", "Двосторонній: новий підхід може виявитися й гіршим, і це теж важливо виявити."),
    ])

    h(doc, "3. Метрики")
    para(doc, "Усі метрики розраховуються поденно для кожної групи та агреговано за весь період.")
    table(doc, ["Тип", "Метрика", "Формула", "Навіщо"], [
        ["Основна (primary)", "CPA — вартість покупки", "Spend / # of Purchase",
         "Головний критерій рішення: чи дешевше ми купуємо покупця"],
        ["Основна (primary)", "Покупки на день", "# of Purchase", "Чи не втрачаємо обсяг продажів"],
        ["Додаткова", "Конверсія клік → покупка", "Purchase / Website Clicks", "Якість трафіку, який приводить реклама"],
        ["Додаткова", "CTR", "Website Clicks / Impressions", "Привабливість нового розміщення/креативів"],
        ["Додаткова", "CPC", "Spend / Website Clicks", "Вартість трафіку"],
        ["Додаткова", "Покрокові конверсії воронки", "Searches / Clicks, View Content / Searches, Add to Cart / View Content, Purchase / Add to Cart",
         "Де саме нова аудиторія поводиться інакше"],
        ["Guardrail", "Витрати на день", "Spend", "Бюджети груп мають бути порівнянними"],
        ["Guardrail", "Охоплення та частота", "Reach; Impressions / Reach", "Не «випалювати» аудиторію надмірною частотою"],
        ["Guardrail", "Повернення / скасування замовлень", "з CRM (поза датасетом)", "Щоб зростання покупок не супроводжувалося зростанням повернень"],
    ], widths=[2.6, 3.6, 5.0, 5.4], align_right_from=None)

    h(doc, "4. Дизайн експерименту")
    kv_table(doc, [
        ("Тип", "Рандомізований контрольований експеримент (A/B-тест) із двома паралельними рекламними кампаніями."),
        ("Одиниця рандомізації", "Користувач рекламної платформи (анонімний ідентифікатор / cookie). Аудиторію випадково ділимо навпіл "
                                 "на рівні рекламного кабінету (split test), щоб той самий користувач не бачив обидва варіанти."),
        ("Групи", "Control Campaign — чинний підхід до розміщення (50 %). Test Campaign — новий підхід (50 %). "
                  "Креативи, продукт, посадкові сторінки та ціни однакові; відрізняється лише спосіб розміщення реклами."),
        ("Бюджет", "Однаковий добовий бюджет для обох кампаній; фактичні витрати фіксуються щодня та порівнюються як guardrail."),
        ("Тривалість", f"30 календарних днів (01.08–30.08.2019): щонайменше чотири повні тижневі цикли, щоб згладити ефект дня тижня "
                       f"та «ефект новизни». Проміжні зупинки за результатами (peeking) не допускаються."),
        ("Період аналізу", "Увесь період без «розігріву»; окремо перевіряємо стабільність результату за тижнями."),
        ("Дані", "Поденна вивантага з рекламного кабінету (Spend, Impressions, Reach, Clicks) та з веб-аналітики "
                 "(Searches, View Content, Add to Cart, Purchase) за трекінг-планом (Завдання №2)."),
    ])

    h(doc, "5. Розмір вибірки та потужність")
    para(doc, f"Базова конверсія клік → покупка в чинній кампанії ≈ {BASE_CR:.0%}. Мінімальний ефект, який хочемо виявити "
              f"(MDE), — відносне зростання на {MDE_REL:.0%} (до {BASE_CR*(1+MDE_REL):.1%}). За α = 0.05 (двосторонній тест) "
              f"та потужності 80 % потрібно приблизно {fmt_int(n_clicks_needed)} кліків у кожній групі. За поточного трафіку "
              f"≈ {fmt_int(clicks_per_day)} кліків на день це досягається вже за ~{days_needed} дні.")
    para(doc, f"Проте основні метрики (CPA, покупки на день) порівнюються на рівні днів, де одиниця спостереження — день "
              f"(n = 30 на групу). За спостережуваної мінливості CPA (SD ≈ ${sd_cpa:.2f} на день) мінімальна виявна різниця CPA "
              f"становить ≈ ${mde_cpa:.2f} (≈ {mde_cpa / C.cpa.mean():.0%} від середнього CPA). Тому тривалість 30 днів — "
              f"розумний компроміс: достатньо для метрик на рівні користувачів і прийнятно для поденних порівнянь; "
              f"якщо різниця виявиться меншою, експеримент продовжують.")

    h(doc, "6. Статистичний метод")
    bullets(doc, [
        ("Поденні метрики (CPA, покупки, конверсії): ", "t-тест Уелча для незалежних вибірок (scipy.stats.ttest_ind, equal_var=False) "
         "з 95 % довірчим інтервалом різниці; як перевірку стійкості — непараметричний критерій Манна–Уїтні, оскільки за n = 30 "
         "розподіл може бути ненормальним."),
        ("Конверсії на рівні користувачів (CTR, клік → покупка, кошик → покупка): ", "z-тест для різниці двох часток "
         "та довірчий інтервал різниці (Вальда/Ньюкомба)."),
        ("Рівень значущості: ", "α = 0.05, двосторонні тести. Для додаткових метрик застосовуємо поправку Бонферроні–Холма, "
         "щоб уникнути хибних відкриттів при множинних порівняннях."),
        ("Перевірки коректності: ", "Sample Ratio Mismatch (частка аудиторії у групах), повнота даних за кожен день, "
         "відсутність фізично неможливих значень (наприклад, покупок більше, ніж додавань у кошик)."),
    ])

    h(doc, "7. Критерії прийняття рішення")
    table(doc, ["Результат", "Рішення"], [
        ["CPA статистично значуще нижчий, покупки не нижчі, guardrails у нормі", "Впроваджуємо новий підхід; плануємо поступове перенесення бюджету"],
        ["CPA не відрізняється, покупки не відрізняються", "Не впроваджуємо; аналізуємо воронку, формулюємо нову гіпотезу (ітерація)"],
        ["CPA значуще вищий або покупки значуще нижчі", "Відхиляємо новий підхід, зупиняємо тестову кампанію"],
        ["Порушені guardrails (нерівні витрати, зламаний трекінг)", "Результат не інтерпретуємо; виправляємо та повторюємо тест"],
    ], widths=[8.6, 8.0], align_right_from=None)

    h(doc, "8. Ризики, припущення та обмеження")
    bullets(doc, [
        ("Ефект новизни: ", "нові розміщення можуть давати тимчасово вищий CTR; тому аналізуємо динаміку за тижнями."),
        ("Сезонність і зовнішні події: ", "обидві групи працюють одночасно, тож зовнішні фактори впливають на них однаково."),
        ("Атрибуція: ", "покупки після кліку можуть відбуватися через кілька днів; використовуємо єдине вікно атрибуції для обох груп."),
        ("Витік між групами: ", "один користувач може побачити обидві кампанії на різних пристроях; мінімізуємо через розділення аудиторії на рівні платформи."),
        ("Якість даних: ", "пропуски за окремі дні та аномалії трекінгу можуть зміщувати поденні метрики — такі дні документуємо й аналізуємо з ними та без них."),
        ("Нерівні витрати: ", "якщо фактичні витрати груп відрізняються, порівнюємо не абсолютні обсяги, а відносні метрики (CPA, конверсії)."),
    ])

    h(doc, "9. Ролі та план робіт")
    table(doc, ["Етап", "Відповідальний", "Строк"], [
        ["Узгодження гіпотези, метрик, бюджету", "Продуктовий аналітик, маркетинг-менеджер", "Тиждень 0"],
        ["Налаштування кампаній і трекінгу (Завдання №2), тестовий запуск", "Маркетинг, веб-аналітик, розробка", "Тиждень 0"],
        ["Проведення експерименту, щоденний моніторинг guardrails", "Маркетинг, аналітик", "Тижні 1–4 (01.08–30.08)"],
        ["Статистичний аналіз і візуалізація (Завдання №3–5)", "Продуктовий аналітик", "Тиждень 5"],
        ["Презентація результатів, рішення", "Аналітик, керівник продукту", "Тиждень 5"],
    ], widths=[7.2, 6.0, 3.4], align_right_from=None)

    out = HERE / f"{SURNAME}_Final Project-1.docx"; doc.save(out); print("written", out)


# ============================================================================= TASK 2 — TRACKING PLAN (PDF)
EVENTS = [
    ("ad_impression", "Показ рекламного оголошення користувачу (рекламна платформа)",
     "campaign_id, campaign_name, experiment_group (Control/Test), placement, creative_id, ad_user_id, device, geo, cost_micros"),
    ("ad_click", "Клік на оголошення з переходом на сайт",
     "campaign_id, experiment_group, placement, creative_id, ad_user_id, click_id (gclid/fbclid), utm_source, utm_medium, utm_campaign, device, geo"),
    ("session_start / page_view", "Початок сесії та перегляд посадкової сторінки після переходу",
     "user_pseudo_id, session_id, click_id, experiment_group, landing_page_url, referrer, device, browser, is_new_user"),
    ("search", "Користувач скористався пошуком на сайті",
     "user_pseudo_id, session_id, experiment_group, search_term, results_count, filters_applied"),
    ("view_content (view_item)", "Перегляд картки продукту / контенту",
     "user_pseudo_id, session_id, experiment_group, item_id, item_name, item_category, price, currency, source (search/catalog/recommendation)"),
    ("add_to_cart", "Додавання продукту до кошика",
     "user_pseudo_id, session_id, experiment_group, item_id, item_category, price, quantity, cart_value, cart_items_count"),
    ("begin_checkout", "Початок оформлення замовлення (додатковий крок для діагностики втрат між кошиком і покупкою)",
     "user_pseudo_id, session_id, experiment_group, cart_value, cart_items_count, payment_method_selected"),
    ("purchase", "Завершена покупка",
     "user_pseudo_id, session_id, experiment_group, order_id, revenue, currency, items (id, qty, price), discount, payment_method, is_first_purchase"),
]
COMMON_ATTRS = [
    ("event_timestamp", "Час події (UTC) — для поденної агрегації та аналізу часу до покупки"),
    ("user_pseudo_id", "Анонімний ідентифікатор користувача (cookie / device id) — унікальні користувачі на кроці"),
    ("session_id", "Ідентифікатор сесії — склеювання кроків воронки в межах візиту"),
    ("experiment_id, experiment_group", "Ідентифікатор експерименту та група (Control / Test) — ключ порівняння"),
    ("click_id / utm_*", "Зв'язок сесії з рекламним кліком і кампанією (атрибуція)"),
    ("device, platform, browser", "Сегментація (мобільні / десктоп), пошук технічних проблем"),
    ("geo (country, region)", "Сегментація за географією"),
    ("page_url, referrer", "Контекст події на сайті"),
]
METRICS_MAP = [
    ("Spend", "Сума cost за ad_impression / дані рекламного кабінету", "Вартість, витрачена за день"),
    ("# of Impressions", "count(ad_impression)", "Покази"),
    ("Reach", "count(distinct ad_user_id) серед ad_impression", "Унікальні охоплені користувачі"),
    ("# of Website Clicks", "count(ad_click)", "Кліки на оголошення"),
    ("# of Searches", "count(distinct user_pseudo_id) з подією search", "Користувачі, що шукали"),
    ("# of View Content", "count(distinct user_pseudo_id) з подією view_content", "Користувачі, що переглядали контент"),
    ("# of Add to Cart", "count(distinct user_pseudo_id) з подією add_to_cart", "Користувачі, що додали в кошик"),
    ("# of Purchase", "count(purchase) або count(distinct order_id)", "Покупки"),
]


def build_task2():
    def rows(items, cols):
        return "\n".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in items) if cols else ""
    html = f"""<!doctype html><html lang="uk"><head><meta charset="utf-8"><title>Трекінг-план</title>
<style>
@page {{ size: A4 landscape; margin: 14mm 14mm 16mm 14mm; }}
body {{ font-family: Arial, Helvetica, sans-serif; color: #0b0b0b; font-size: 10.5pt; line-height: 1.35; margin: 0; }}
h1 {{ font-size: 20pt; margin: 0 0 4px; }} h2 {{ font-size: 13pt; color: #1c5cab; margin: 18px 0 6px; }}
.sub {{ color: #52514e; font-size: 10pt; margin-bottom: 4px; }}
.meta {{ color: #52514e; font-size: 9pt; border-bottom: 1px solid #c3c2b7; padding-bottom: 8px; margin-bottom: 10px; }}
table {{ border-collapse: collapse; width: 100%; page-break-inside: auto; }}
tr {{ page-break-inside: avoid; }}
th, td {{ border: 1px solid #c3c2b7; padding: 5px 7px; vertical-align: top; text-align: left; font-size: 9.5pt; }}
th {{ background: #e8eef7; }} tr:nth-child(even) td {{ background: #f7f7f4; }}
code {{ font-family: Menlo, Consolas, monospace; font-size: 9pt; background: #f1f0ea; padding: 0 3px; border-radius: 3px; }}
ul {{ margin: 4px 0 8px 18px; padding: 0; }} li {{ margin-bottom: 3px; }}
.pb {{ page-break-before: always; }}
</style></head><body>
<h1>Трекінг-план A/B-експерименту</h1>
<div class="sub">Новий підхід до розміщення онлайн-реклами продукту · Фінальний проєкт, Завдання №2</div>
<div class="meta">Автор: {AUTHOR} · Дисципліна: Продуктова аналітика та прикладна статистика (GoIT Neoversity) · {DATE}</div>

<h2>1. Мета трекінгу</h2>
<p>Зафіксувати детальну поведінку користувачів від показу реклами до покупки так, щоб для контрольної (<code>Control</code>) і тестової
(<code>Test</code>) груп можна було поденно порівнювати ефективність на рівні кожної події воронки: показ → клік → сесія на сайті →
пошук → перегляд контенту → додавання в кошик → оформлення → покупка. Кожна подія має нести ознаку групи експерименту та ідентифікатори,
що дозволяють склеїти рекламний клік із діями на сайті.</p>

<h2>2. Воронка та події</h2>
<table><tr><th style="width:17%">Подія</th><th style="width:30%">Коли спрацьовує</th><th>Атрибути події</th></tr>
{rows(EVENTS, True)}
</table>

<h2>3. Спільні атрибути для всіх подій</h2>
<table><tr><th style="width:28%">Атрибут</th><th>Призначення</th></tr>
{rows(COMMON_ATTRS, True)}
</table>

<h2>4. Відповідність подій полям датасету експерименту</h2>
<table><tr><th style="width:22%">Поле датасету</th><th style="width:43%">Розрахунок з подій (за день і групу)</th><th>Зміст</th></tr>
{rows(METRICS_MAP, True)}
</table>

<h2>5. Похідні метрики для порівняння груп</h2>
<ul>
<li><b>CTR</b> = # of Website Clicks / # of Impressions; <b>CPC</b> = Spend / # of Website Clicks; <b>Частота</b> = Impressions / Reach.</li>
<li><b>Покрокові конверсії</b>: Searches / Clicks, View Content / Searches, Add to Cart / View Content, Purchase / Add to Cart.</li>
<li><b>Наскрізна конверсія</b> = Purchase / Clicks; <b>CPA</b> = Spend / Purchase; <b>ROAS</b> = Revenue / Spend (потребує <code>revenue</code> у події <code>purchase</code>).</li>
<li><b>Якість аудиторії</b>: частка нових користувачів, середня кількість переглянутих товарів, час від кліку до покупки.</li>
</ul>

<h2>6. Правила збору та контролю якості</h2>
<ul>
<li>Група експерименту (<code>experiment_group</code>) присвоюється на рівні рекламної платформи й передається на сайт через <code>click_id</code>/<code>utm_campaign</code>; на сайті вона зберігається в cookie на весь період тесту, щоб усі наступні події користувача мали ту саму групу.</li>
<li>Кожна подія надсилається один раз (дедуплікація за <code>event_id</code>); <code>purchase</code> підтверджується серверно (server-side) за <code>order_id</code>, щоб уникнути дублів і блокувань скриптів.</li>
<li>Агрегація за день виконується в єдиному часовому поясі (UTC) для рекламних даних і веб-аналітики.</li>
<li>Щоденні автоматичні перевірки: повнота даних за день для обох груп; відсутність фізично неможливих значень (наприклад, Purchase &gt; Add to Cart); частка аудиторії у групах близька до 50/50 (Sample Ratio Mismatch); узгодженість кліків у рекламному кабінеті та сесій на сайті (±10 %).</li>
<li>Персональні дані (e-mail, телефон, адреса) у події не передаються — лише анонімні ідентифікатори.</li>
</ul>

<h2>7. Інструменти</h2>
<p>Рекламні події — рекламний кабінет (Meta Ads / Google Ads) з експортом поденної статистики; події сайту — Google Analytics 4 / Amplitude
через Google Tag Manager з увімкненим server-side підтвердженням покупок; сховище — BigQuery, щоденна вітрина «день × група × метрики»,
з якої будуються звіти та розрахунки Завдань №3–5.</p>
</body></html>"""
    html_path = FIG / "tracking_plan.html"; html_path.write_text(html, encoding="utf-8")
    out = HERE / f"{SURNAME}_Final Project-2.pdf"
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={out}", html_path.as_uri()], check=True, capture_output=True, timeout=120)
    print("written", out)


# ============================================================================= TASK 3 — RESULTS (DOCX)
def build_task3():
    doc = new_doc("Результати A/B-тесту та їх статистична значущість",
                  "Новий підхід до розміщення онлайн-реклами продукту · Фінальний проєкт, Завдання №3")

    h(doc, "1. Коротко про результат")
    para(doc, f"Новий підхід до розміщення реклами не збільшив кількість покупок: у тестовій групі в середньому "
              f"{tests.loc['purchase','Test']:.0f} покупок на день проти {tests.loc['purchase','Control']:.0f} у контрольній "
              f"({ptxt(tests.loc['purchase','p_welch'])}). Водночас витрати тестової групи були на "
              f"{tests.loc['spend','rel']:+.0%} вищими ({ptxt(tests.loc['spend','p_welch'])}), тому вартість покупки (CPA) "
              f"зросла з ${tests.loc['cpa','Control']:.2f} до ${tests.loc['cpa','Test']:.2f} на день "
              f"({tests.loc['cpa','rel']:+.0%}, {ptxt(tests.loc['cpa','p_welch'])} — різниця статистично незначуща). "
              f"Новий підхід істотно змінив структуру воронки: CTR зріс удвічі, конверсія клік → кошик упала, а конверсія "
              f"кошик → покупка зросла. Основну гіпотезу (нижчий CPA) не підтверджено.")

    h(doc, "2. Дані та підготовка")
    bullets(doc, [
        ("Джерело: ", "поденні дані двох кампаній за 01.08–30.08.2019 (30 днів): витрати, покази, охоплення та кроки воронки сайту."),
        ("Пропуски: ", "у контрольній групі за 05.08 є лише витрати, показники воронки відсутні; щоб порівнювати групи за "
                       "однаковий період, цей день виключено з розрахунків для обох груп (по 29 днів у кожній)."),
        ("Аномалії: ", "у контрольній групі " + ", ".join(f"{r.date:%d.%m}" for r in anom.itertuples()) +
         " кількість покупок перевищує кількість додавань у кошик (конверсія кроку > 100 %) — імовірна похибка трекінгу. "
         "День залишено в основному аналізі; чутливість результату до нього перевірено окремо (Завдання №5)."),
        ("Одиниця аналізу: ", "день (n = 29 у кожній групі) для абсолютних показників та середніх конверсій; користувач/подія — для "
                              "агрегованих часток (z-тест пропорцій)."),
    ])

    h(doc, "3. Метод")
    para(doc, "Основний метод — t-тест Уелча для незалежних вибірок (scipy.stats.ttest_ind, equal_var=False): він не вимагає "
              "рівності дисперсій і стійкий до відмінностей у мінливості метрик між групами. Для кожної метрики наведено 95 % довірчий "
              "інтервал різниці середніх. Оскільки за n ≈ 30 розподіли деяких метрик відрізняються від нормального (тест "
              "Шапіро–Уїлка відхиляє нормальність для кліків, покупок, CPA та окремих конверсій), результати перевірено "
              "непараметричним критерієм Манна–Уїтні. Для конверсій на рівні користувачів (CTR, клік → кошик, кошик → покупка, "
              "клік → покупка) додатково застосовано z-тест для різниці двох часток. Рівень значущості α = 0.05, тести двосторонні.")

    h(doc, "4. Підсумкові показники за період")
    rows = [[METRIC_UA["spend"].replace(" / день", " (сума)"), fmt_money(tot.loc["spend", "Control"]), fmt_money(tot.loc["spend", "Test"]),
             f"{tot.loc['spend','Test']/tot.loc['spend','Control']-1:+.1%}"]]
    for s in STEPS:
        rows.append([STEP_UA[s], fmt_int(tot.loc[s, "Control"]), fmt_int(tot.loc[s, "Test"]), f"{tot.loc[s,'Test']/tot.loc[s,'Control']-1:+.1%}"])
    for m in ["ctr", "cpc", "cart_rate", "cr_cart", "cr_click", "cpa"]:
        f = (lambda v: f"${v:.2f}") if m in ("cpc", "cpa") else (lambda v: fmt_pct(v, 2))
        rows.append([METRIC_UA[m], f(ratios.loc[m, "Control"]), f(ratios.loc[m, "Test"]), f"{ratios.loc[m,'Test']/ratios.loc[m,'Control']-1:+.1%}"])
    table(doc, ["Показник (29 днів з повними даними в обох групах)", "Control", "Test", "Test vs Control"], rows, widths=[7.0, 3.2, 3.2, 3.2])

    h(doc, "5. Порівняння поденних метрик: t-тест Уелча та критерій Манна–Уїтні")
    rows = []
    for m in ["spend", "impressions", "reach", "clicks", "add_to_cart", "purchase", "ctr", "cpc", "cart_rate", "cr_cart", "cr_click", "cpa"]:
        r = tests.loc[m]
        if m in ("ctr", "cart_rate", "cr_cart", "cr_click"):
            fv = lambda v: fmt_pct(v, 1); fd = fmt_pp(r["diff"], 1); fci = f"[{fmt_pp(r.ci_lo,1)}; {fmt_pp(r.ci_hi,1)}]"
        elif m in ("cpc", "cpa"):
            fv = lambda v: f"${v:.2f}"; fd = f"${r['diff']:+.2f}"; fci = f"[{r.ci_lo:+.2f}; {r.ci_hi:+.2f}]"
        else:
            fv = lambda v: fmt_int(v); fd = f"{r['diff']:+,.0f}".replace(",", " "); fci = f"[{r.ci_lo:+,.0f}; {r.ci_hi:+,.0f}]".replace(",", " ")
        sig = "так" if r.p_welch < 0.05 else "ні"
        rows.append([METRIC_UA[m], fv(r.Control), fv(r.Test), fd, fci, f"{r.t:.2f}", fmt_p(r.p_welch), fmt_p(r.p_mw), sig])
    table(doc, ["Метрика", "Control", "Test", "Різниця", "95% ДІ", "t", "p (Уелч)", "p (М–У)", "Значуща"], rows,
          widths=[3.6, 1.7, 1.7, 1.7, 3.0, 1.1, 1.3, 1.3, 1.2], font_size=8)

    doc.add_picture(str(P_EFFECTS), width=Cm(16.6))
    para(doc, "Рис. 1. Відносна різниця Test vs Control для ключових метрик (середні по днях) з 95 % довірчими інтервалами.",
         italic=True, size=9)

    h(doc, "6. Конверсії на рівні користувачів: z-тест для різниці часток")
    rows = []
    for name, r in [("CTR (кліки / покази)", prop_ctr), ("Клік → кошик", prop_cartrate), ("Кошик → покупка", prop_cart), ("Клік → покупка", prop_click)]:
        rows.append([name, fmt_pct(r["p2"], 2), fmt_pct(r["p1"], 2), fmt_pp(r["diff"], 2), f"[{fmt_pp(r['lo'],2)}; {fmt_pp(r['hi'],2)}]",
                     f"{r['z']:.1f}", fmt_p(r["p"])])
    table(doc, ["Конверсія", "Control", "Test", "Різниця", "95% ДІ (Вальд)", "z", "p"], rows, widths=[4.2, 2.1, 2.1, 2.3, 3.6, 1.2, 1.3], font_size=8.5)
    para(doc, "Агреговані тести мають величезну потужність (десятки й сотні тисяч подій), тому навіть невеликі різниці "
              "виходять «значущими». Вони припускають незалежність подій і однорідність днів, чого на практиці немає, тому "
              "для рішень ми спираємося передусім на поденні тести (розділ 5), а агреговані використовуємо як підтвердження напряму ефекту.",
         italic=True, size=9.5)

    doc.add_picture(str(P_DAILY), width=Cm(16.6))
    para(doc, "Рис. 2. Щоденна динаміка покупок, CPA та конверсії кошик → покупка за групами.", italic=True, size=9)

    h(doc, "7. Інтерпретація результатів")
    bullets(doc, [
        ("Основна метрика — CPA — не покращилась. ", f"Середній денний CPA у тестовій групі ${tests.loc['cpa','Test']:.2f} проти "
         f"${tests.loc['cpa','Control']:.2f} у контрольній; різниця {tests.loc['cpa','rel']:+.0%} статистично незначуща "
         f"({ptxt(tests.loc['cpa','p_welch'])}), довірчий інтервал включає нуль. Нульову гіпотезу про відсутність різниці "
         "відхилити не можна; при цьому напрям ефекту — не на користь нового підходу."),
        ("Кількість покупок не змінилася. ", f"{tests.loc['purchase','Test']:.0f} проти {tests.loc['purchase','Control']:.0f} покупок на день "
         f"({ptxt(tests.loc['purchase','p_welch'])}) — практично ідентичний результат за вищих витрат "
         f"({tests.loc['spend','rel']:+.0%}, {ptxt(tests.loc['spend','p_welch'])})."),
        ("Верх воронки став ефективнішим. ", f"Показів і охоплення у тестовій групі значно менше ({tests.loc['impressions','rel']:+.0%} та "
         f"{tests.loc['reach','rel']:+.0%}, p < 0.001), але CTR зріс удвічі ({fmt_pct(tests.loc['ctr','Control'])} → "
         f"{fmt_pct(tests.loc['ctr','Test'])}, {ptxt(tests.loc['ctr','p_welch'])}); кліків стало більше за майже тієї самої вартості кліку."),
        ("Середина воронки просіла, низ — зріс. ", f"Конверсія клік → кошик упала з {fmt_pct(tests.loc['cart_rate','Control'])} до "
         f"{fmt_pct(tests.loc['cart_rate','Test'])} ({ptxt(tests.loc['cart_rate','p_welch'])}), тоді як конверсія кошик → покупка "
         f"зросла з {fmt_pct(tests.loc['cr_cart','Control'])} до {fmt_pct(tests.loc['cr_cart','Test'])} ({ptxt(tests.loc['cr_cart','p_welch'])}). "
         "Нова аудиторія частіше клікає й рішучіше купує, якщо знаходить потрібний товар, але рідше доходить до кошика — "
         "ймовірно, через невідповідність очікувань від оголошення та вмісту сайту."),
        ("Наскрізна конверсія клік → покупка ", f"нижча у тестовій групі ({fmt_pct(tests.loc['cr_click','Control'])} → "
         f"{fmt_pct(tests.loc['cr_click','Test'])}), але поденно різниця незначуща ({ptxt(tests.loc['cr_click','p_welch'])})."),
        ("Стійкість висновків. ", "Непараметричний критерій Манна–Уїтні дає ті самі висновки щодо значущості для всіх метрик; "
         "виключення аномального дня 30.08 не змінює напряму ефектів."),
    ])

    h(doc, "8. Висновки та рекомендації")
    bullets(doc, [
        "Гіпотезу про те, що новий підхід до розміщення реклами знижує вартість покупки, не підтверджено: покупок стільки ж, витрати вищі, CPA не кращий. У поточному вигляді новий підхід впроваджувати не варто.",
        "Новий підхід має сильні сторони — удвічі вищий CTR і значно вищу готовність до покупки після додавання в кошик. Проблема локалізована на кроці клік → кошик. Рекомендується доопрацювати посадкові сторінки / каталог під очікування нової аудиторії та повторити тест.",
        "Якщо вдасться підняти конверсію клік → кошик до рівня контрольної групи, поєднання вищого CTR і вищої конверсії кошик → покупка дасть істотний приріст покупок за той самий бюджет — це головна гіпотеза для наступної ітерації.",
        "Перед наступним запуском слід усунути проблеми трекінгу (пропущений день, покупки без кроку «кошик») та забезпечити однакові витрати груп, інакше порівняння абсолютних показників буде некоректним.",
    ], style="List Number")

    out = HERE / f"{SURNAME}_Final Project-3.docx"; doc.save(out); print("written", out)


if __name__ == "__main__":
    build_task1(); build_task2(); build_task3()
