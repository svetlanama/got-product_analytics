# -*- coding: utf-8 -*-
"""Build Moiseyenko_Svitlana_PA_assignment_6.xlsx: cohort LTV analysis of the
Superstore dataset (cohort = quarter of the customer's first order, cumulative
LTV by year) with heatmaps. Needs pandas, xlrd, openpyxl."""
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "Sample - Superstore.xls")
OUT = os.path.join(HERE, "Moiseyenko_Svitlana_PA_assignment_6.xlsx")

# ----------------------------------------------------------------- palette
BLUE, GRID, MUTED, INK2 = "2A78D6", "E1E0D9", "898781", "52514E"
SEQ_LO, SEQ_HI = "EAF3FD", "4A90E2"
HEAD_FILL = PatternFill("solid", fgColor="F1F0EA")
THIN = Side(style="thin", color=GRID)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="top", wrap_text=True)

# ----------------------------------------------------------------- data
src = pd.ExcelFile(SRC)
orders, people, returns = (src.parse(s) for s in ("Orders", "People", "Returns"))
N = len(orders)

first = orders.groupby("Customer ID")["Order Date"].min()
o = orders.join(first.rename("first"), on="Customer ID")
o["coh"] = o["first"].dt.year.astype(str) + " Q" + o["first"].dt.quarter.astype(str)
o["yr"] = o["Order Date"].dt.year
o["life"] = o["yr"] - o["first"].dt.year + 1
YEARS = sorted(o["yr"].unique())
COHORTS = sorted(o["coh"].unique())
size = o.groupby("coh")["Customer ID"].nunique()
rev = o.pivot_table(index="coh", columns="yr", values="Sales", aggfunc="sum").reindex(COHORTS)
ltv = rev.fillna(0).cumsum(axis=1).div(size, axis=0)          # by calendar year
ltv = ltv.where(rev.notna())
rev_l = o.pivot_table(index="coh", columns="life", values="Sales", aggfunc="sum").reindex(COHORTS)
ltv_l = rev_l.fillna(0).cumsum(axis=1).div(size, axis=0).where(rev_l.notna())  # by year of life

customers = (orders.groupby("Customer ID")["Customer Name"].first().reset_index()
             .sort_values("Customer ID").reset_index(drop=True))
M = len(customers)


def blend(v, lo, hi):
    """Static heatmap colour (same ramp as the conditional colour scale)."""
    t = 0 if hi == lo else (v - lo) / (hi - lo)
    ch = [round(int(SEQ_LO[i:i + 2], 16) + t * (int(SEQ_HI[i:i + 2], 16) - int(SEQ_LO[i:i + 2], 16)))
          for i in (0, 2, 4)]
    return "%02X%02X%02X" % tuple(ch)


def header(ws, row, col, labels):
    for j, text in enumerate(labels):
        c = ws.cell(row, col + j, text)
        c.font, c.fill, c.alignment, c.border = Font(bold=True), HEAD_FILL, CENTER, BOX


# ----------------------------------------------------------------- workbook
wb = Workbook()
ws = wb.active
ws.title = "Когортний LTV"
wc = wb.create_sheet("Customers")
wo = wb.create_sheet("Orders")

# --- Orders: source data + helper columns (live formulas)
cols = list(orders.columns)
C = {name: get_column_letter(i + 1) for i, name in enumerate(cols)}
H1, H2, H3, H4 = (get_column_letter(len(cols) + k) for k in (1, 2, 3, 4))
wo.append(cols + ["Перша покупка", "Когорта", "Рік замовлення", "Рік життя"])
for r, row in enumerate(orders.itertuples(index=False), start=2):
    wo.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else v for v in row] + [
        f"=INDEX(Customers!$C$2:$C${M + 1},MATCH({C['Customer ID']}{r},Customers!$A$2:$A${M + 1},0))",
        f"=INDEX(Customers!$D$2:$D${M + 1},MATCH({C['Customer ID']}{r},Customers!$A$2:$A${M + 1},0))",
        f"=YEAR({C['Order Date']}{r})",
        f"={H3}{r}-YEAR({H1}{r})+1",
    ])
    for col in (C["Order Date"], C["Ship Date"], H1):
        wo[f"{col}{r}"].number_format = "yyyy-mm-dd"
for c in wo[1]:
    c.font, c.fill = Font(bold=True), HEAD_FILL
for i in range(1, len(cols) + 5):
    wo.column_dimensions[get_column_letter(i)].width = 15
wo.freeze_panes = "A2"

for name, df in (("People", people), ("Returns", returns)):
    w = wb.create_sheet(name)
    w.append(list(df.columns))
    for row in df.itertuples(index=False):
        w.append(list(row))
    for c in w[1]:
        c.font, c.fill = Font(bold=True), HEAD_FILL
    for i in range(1, len(df.columns) + 1):
        w.column_dimensions[get_column_letter(i)].width = 20

# --- Customers: one row per customer, first order date and cohort
O_DATE = f"Orders!${C['Order Date']}$2:${C['Order Date']}${N + 1}"
O_CUST = f"Orders!${C['Customer ID']}$2:${C['Customer ID']}${N + 1}"
O_SALES = f"Orders!${C['Sales']}$2:${C['Sales']}${N + 1}"
O_COH = f"Orders!${H2}$2:${H2}${N + 1}"
O_YEAR = f"Orders!${H3}$2:${H3}${N + 1}"
O_LIFE = f"Orders!${H4}$2:${H4}${N + 1}"
header(wc, 1, 1, ["Customer ID", "Customer Name", "Перша покупка", "Когорта (квартал залучення)", "Рік залучення"])
for r, (cid, cname) in enumerate(customers.itertuples(index=False), start=2):
    wc.append([cid, cname, f"=MINIFS({O_DATE},{O_CUST},A{r})",
               f'=YEAR(C{r})&" Q"&ROUNDUP(MONTH(C{r})/3,0)', f"=YEAR(C{r})"])
    wc[f"C{r}"].number_format = "yyyy-mm-dd"
for col, w in zip("ABCDE", (14, 24, 16, 18, 14)):
    wc.column_dimensions[col].width = w
wc.row_dimensions[1].height = 32
wc.freeze_panes = "A2"

# --- Когортний LTV: summary sheet
K = len(COHORTS)
ny = len(YEARS)
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 12
for i in range(3, 3 + ny + 1):
    ws.column_dimensions[get_column_letter(i)].width = 13
ws.column_dimensions[get_column_letter(4 + ny)].width = 3
ws.column_dimensions[get_column_letter(5 + ny)].width = 90

ws["A1"] = "Когортний аналіз LTV — Superstore"
ws["A1"].font = Font(bold=True, size=16, color=BLUE)
ws["A2"] = ("Когорта — квартал першої покупки клієнта. LTV — кумулятивний дохід (Sales) когорти, "
            "поділений на кількість клієнтів у когорті, $ на клієнта.")
ws["A2"].font = Font(color=INK2)


def title(row, text, note=None):
    ws.cell(row, 1, text).font = Font(bold=True, size=12)
    if note:
        ws.cell(row + 1, 1, note).font = Font(italic=True, color=MUTED)


def cohort_rows(top):
    """Cohort label + size columns; returns first data row."""
    for i, coh in enumerate(COHORTS):
        r = top + i
        a = ws.cell(r, 1, coh)
        a.font, a.border, a.alignment = Font(bold=True), BOX, CENTER
        b = ws.cell(r, 2, f"=COUNTIF(Customers!$D$2:$D${M + 1},$A{r})")
        b.border, b.alignment = BOX, CENTER
    return top


# Table 1: revenue by cohort and order year
T1 = 6
title(T1 - 2, "1. Дохід когорти за роками замовлення, $")
header(ws, T1 - 1, 1, ["Когорта", "Клієнтів"] + YEARS + ["Разом"])
cohort_rows(T1)
for i in range(K):
    r = T1 + i
    for j in range(ny):
        col = get_column_letter(3 + j)
        c = ws.cell(r, 3 + j, f"=SUMIFS({O_SALES},{O_COH},$A{r},{O_YEAR},{col}${T1 - 1})")
        c.number_format, c.border = "#,##0", BOX
    c = ws.cell(r, 3 + ny, f"=SUM(C{r}:{get_column_letter(2 + ny)}{r})")
    c.number_format, c.border, c.font = "#,##0", BOX, Font(bold=True)
r = T1 + K
ws.cell(r, 1, "Разом").font = Font(bold=True)
for j in range(2, 4 + ny):
    col = get_column_letter(j)
    c = ws.cell(r, j, f"=SUM({col}{T1}:{col}{r - 1})")
    c.font, c.border, c.number_format = Font(bold=True), BOX, "#,##0"
    c.alignment = CENTER if j == 2 else Alignment()
ws.cell(r, 1).border = BOX
ws.cell(r, 1).alignment = CENTER


def heatmap(top, values, formula):
    lo, hi = values.min().min(), values.max().max()
    for i, coh in enumerate(COHORTS):
        r = top + i
        for j, key in enumerate(values.columns):
            c = ws.cell(r, 3 + j, formula(r, j))
            c.number_format, c.border, c.alignment = "#,##0", BOX, CENTER
            v = values.loc[coh, key]
            if pd.notna(v):
                c.fill = PatternFill("solid", fgColor=blend(v, lo, hi))
    rng = f"C{top}:{get_column_letter(2 + len(values.columns))}{top + K - 1}"
    ws.conditional_formatting.add(rng, ColorScaleRule(
        start_type="min", start_color=SEQ_LO, end_type="max", end_color=SEQ_HI))


# Table 2: cumulative LTV by calendar year (heatmap)
T2 = T1 + K + 5
title(T2 - 3, "2. Кумулятивний LTV за календарними роками, $ на клієнта (теплова карта)",
      "Накопичений дохід когорти на кінець року / кількість клієнтів когорти.")
header(ws, T2 - 1, 1, ["Когорта", "Клієнтів"] + YEARS)
cohort_rows(T2)
heatmap(T2, ltv, lambda r, j: (
    f'=IF({get_column_letter(3 + j)}${T2 - 1}<VALUE(LEFT($A{r},4)),"",'
    f"SUM($C{r - T2 + T1}:{get_column_letter(3 + j)}{r - T2 + T1})/$B{r})"))

# Table 3: cumulative LTV by year of cohort life (heatmap)
T3 = T2 + K + 4
title(T3 - 3, "3. Кумулятивний LTV за роками життя когорти, $ на клієнта (теплова карта)",
      f"Рік 1 — календарний рік залучення, рік 2 — наступний і т. д. Дані закінчуються {YEARS[-1]} роком.")
header(ws, T3 - 1, 1, ["Когорта", "Клієнтів"] + list(range(1, ny + 1)))
cohort_rows(T3)
heatmap(T3, ltv_l, lambda r, j: (
    f'=IF(VALUE(LEFT($A{r},4))+{get_column_letter(3 + j)}${T3 - 1}-1>{YEARS[-1]},"",'
    f'SUMIFS({O_SALES},{O_COH},$A{r},{O_LIFE},"<="&{get_column_letter(3 + j)}${T3 - 1})/$B{r})'))

# Conclusions
y0, yl = YEARS[0], YEARS[-1]
by_year = o.groupby(o["first"].dt.year)["Customer ID"].nunique()
q0 = [c for c in COHORTS if c.startswith(str(y0))]
best = ltv[yl].idxmax()
old_share = rev.loc[q0, yl].sum() / rev[yl].sum()
q1 = [c for c in COHORTS if c.endswith("Q1")]
small = [c for c in COHORTS if size[c] < 10]
notes = [
    ("Висновки", None),
    ("Найцінніші — когорти першого року.",
     f"Когорта {best} має найвищий LTV: ${ltv.loc[best, yl]:,.0f} на клієнта на кінець {yl} року. "
     f"Усі когорти {y0} року накопичили ${ltv.loc[q0, yl].min():,.0f}–${ltv.loc[q0, yl].max():,.0f} за чотири роки, "
     "і серед них що раніше долучився клієнт, то вищий його LTV."),
    ("LTV не виходить на плато.",
     f"Клієнти купують повторно щороку: когорти {y0} року додають приблизно $600–1 000 на клієнта за рік, "
     f"і приріст не згасає: у {yl} році він становив ${(ltv.loc[q0, yl] - ltv.loc[q0, yl - 1]).min():,.0f}–"
     f"${(ltv.loc[q0, yl] - ltv.loc[q0, yl - 1]).max():,.0f}. Утримання працює добре."),
    ("Залучення нових клієнтів різко падає.",
     "Нових клієнтів за роками: " + ", ".join(f"{y} — {n}" for y, n in by_year.items()) +
     f". Когорти {y0} року — це {size[q0].sum() / size.sum():.0%} клієнтської бази, і вони дали "
     f"{old_share:.0%} доходу {yl} року. Зростання тримається на старих клієнтах."),
    ("Новіші когорти стартують слабше.",
     "LTV першого року для когорт Q1 (повний рік спостереження): " +
     ", ".join(f"{c} — ${ltv_l.loc[c, 1]:,.0f}" for c in q1) +
     ". Кожна наступна когорта приносить менше в рік залучення, хоча на другий-третій рік "
     "когорти 2021–2022 частково наздоганяють."),
    ("Обмеження.",
     f"Когорти {', '.join(small)} мають менше 10 клієнтів — їхній LTV нестабільний і не годиться для висновків. "
     "Когорти Q4 мають лише до трьох місяців у першому році, тому в таблиці 3 порівнювати варто когорти одного кварталу. "
     "LTV рахується за доходом (Sales) без урахування повернень і прибутковості."),
    ("Рекомендації.",
     f"Орієнтир для вартості залучення: клієнт приносить близько ${ltv_l[1].loc[q0 + [c for c in COHORTS if c.startswith(str(y0 + 1))]].mean():,.0f} "
     f"доходу в рік залучення і близько ${ltv.loc[q0, yl].mean():,.0f} за чотири роки. Пріоритет — відновити залучення нових клієнтів "
     "і зберегти програми утримання, які вже дають стабільні повторні покупки."),
]
NC = 5 + ny
row = T1 - 2
for head, body in notes:
    if body is None:
        ws.cell(row, NC, head).font = Font(bold=True, size=12)
        row += 2
        continue
    ws.cell(row, NC, head).font = Font(bold=True, color=BLUE)
    c = ws.cell(row + 1, NC, body)
    c.alignment = WRAP
    ws.merge_cells(start_row=row + 1, start_column=NC, end_row=row + 3, end_column=NC)
    row += 5

ws.freeze_panes = "C1"
wb.save(OUT)
print("saved", OUT)
print(pd.concat([size.rename("n"), ltv.round(0), ltv_l.round(0)], axis=1).to_string())
